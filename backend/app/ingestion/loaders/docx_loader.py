"""DOCX document loader using python-docx."""
from pathlib import Path
from typing import List, Tuple

from docx import Document
from docx.enum.text import WD_PARAGRAPH_ALIGNMENT

from app.core.logging import get_logger
from app.core.exceptions import IngestionError

logger = get_logger("ingestion.docx")


def load_docx(file_path: str) -> Tuple[str, List[dict]]:
    """Load DOCX and return full text + paragraph-level metadata.

    Returns:
        Tuple of (full_text, elements_metadata)
    """
    try:
        doc = Document(file_path)
        elements = []
        full_text_parts = []

        for i, para in enumerate(doc.paragraphs):
            if not para.text.strip():
                continue

            style = para.style.name if para.style else "Normal"
            is_heading = style.startswith("Heading") or para.text.strip().startswith("#")

            element = {
                "index": i,
                "text": para.text,
                "style": style,
                "is_heading": is_heading,
                "alignment": str(para.alignment) if para.alignment else None,
            }
            elements.append(element)
            full_text_parts.append(para.text)

        # Extract tables
        for table_idx, table in enumerate(doc.tables):
            table_text = []
            for row in table.rows:
                row_text = [cell.text for cell in row.cells]
                table_text.append(" | ".join(row_text))
            if table_text:
                elements.append({
                    "index": len(elements),
                    "text": "\n".join(table_text),
                    "style": "Table",
                    "is_heading": False,
                    "table_index": table_idx,
                })
                full_text_parts.append("\n".join(table_text))

        full_text = "\n\n".join(full_text_parts)
        logger.info(f"Loaded DOCX: {Path(file_path).name}, {len(elements)} elements")
        return full_text, elements

    except Exception as e:
        raise IngestionError(f"Failed to load DOCX {file_path}: {e}")
